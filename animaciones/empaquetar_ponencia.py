"""Empaqueta el material para PowerPoint y arma montajes por bloque.

Entradas : ../exports/<tema>/{oscuro,claro}/<Pieza>.mp4         (render_todo.sh)
           ../exports/png/{oscuro,claro}/<tema>/<Pieza>.png      (render_fotos.sh)
           ../exports/png/stickers/{oscuro,claro}/<tema>/<Pieza>.png (postproceso_fotos.py)
Salidas  : ../exports/presentaciones/ponencia_completa_{oscuro,claro}.pptx  (65 diapositivas + secciones)
           ../exports/presentaciones/ponencia_seleccion_{oscuro,claro}.pptx (solo las marcadas 'recomendada')
           ../exports/presentaciones/stickers_{oscuro,claro}.pptx           (biblioteca de stickers en cuadrícula)
           ../exports/montajes/NN_<bloque>_{oscuro,claro}.mp4 y reel_recomendadas_*.mp4
Uso: python3 empaquetar_ponencia.py [presentaciones|stickers|montajes|todo]
No se borra ni se filtra nada de exports/: 'seleccion' es un archivo aparte.
"""
import io
import re
import subprocess
import sys
import unicodedata
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).parent))
from catalogo import BLOQUES  # noqa: E402

Image.MAX_IMAGE_PIXELS = None
EXP = Path(__file__).resolve().parent.parent / "exports"
FONDO = {"oscuro": RGBColor(0x0B, 0x1F, 0x3A), "claro": RGBColor(0xFF, 0xFF, 0xFF)}
TINTA = {"oscuro": RGBColor(0xF1, 0xF5, 0xF9), "claro": RGBColor(0x0F, 0x17, 0x2A)}
ACENTO = {"oscuro": RGBColor(0xF5, 0x9E, 0x0B), "claro": RGBColor(0xD9, 0x77, 0x06)}
W, H = Inches(13.333), Inches(7.5)


def slug(texto):
    t = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", t.lower()).strip("_")


def nombre_legible(clase):
    return re.sub(r"(?<=[a-záéíóú0-9])(?=[A-Z])|(?<=[A-Z])(?=[A-Z][a-z])", " ", clase)


def nueva(tema):
    p = Presentation()
    p.slide_width, p.slide_height = W, H
    return p


def diapositiva(prs, tema):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = FONDO[tema]
    return s


def texto(slide, txt, x, y, w, h, tam, color, negrita=False, centrado=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER if centrado else PP_ALIGN.LEFT
    r = p.add_run()
    r.text = txt
    r.font.size, r.font.bold, r.font.name = Pt(tam), negrita, "Carlito"
    r.font.color.rgb = color
    return tb


def poster(tema, carpeta, clase):
    """Fotograma opaco 1920x1080 para el póster del video."""
    src = EXP / "png" / tema / carpeta / f"{clase}.png"
    fondo = tuple(FONDO[tema])
    im = Image.open(src).convert("RGBA")
    im.thumbnail((1920, 1080), Image.LANCZOS)
    lienzo = Image.new("RGB", (1920, 1080), fondo)
    lienzo.paste(im, ((1920 - im.width) // 2, (1080 - im.height) // 2), im)
    b = io.BytesIO()
    lienzo.save(b, "JPEG", quality=90)
    b.seek(0)
    return b


def presentaciones():
    (EXP / "presentaciones").mkdir(parents=True, exist_ok=True)
    for tema in ("oscuro", "claro"):
        for variante in ("completa", "seleccion"):
            prs, n = nueva(tema), 0
            for i, (bloque, carpeta, piezas) in enumerate(BLOQUES, 1):
                elegidas = [x for x in piezas if variante == "completa" or x[3]]
                if not elegidas:
                    continue
                s = diapositiva(prs, tema)
                texto(s, f"{i:02d}", Inches(0.9), Inches(2.4), Inches(4), Inches(1.2), 54, ACENTO[tema], True)
                texto(s, bloque, Inches(0.9), Inches(3.5), Inches(11.5), Inches(1.2), 44, TINTA[tema], True)
                for clase, muestra, uso, rec in elegidas:
                    mp4 = EXP / carpeta / tema / f"{clase}.mp4"
                    if not mp4.exists():
                        print("falta", mp4)
                        continue
                    s = diapositiva(prs, tema)
                    s.shapes.add_movie(str(mp4), 0, 0, W, H, poster_frame_image=poster(tema, carpeta, clase),
                                       mime_type="video/mp4")
                    s.notes_slide.notes_text_frame.text = (
                        f"{nombre_legible(clase)}{'  [recomendada]' if rec else ''}\n"
                        f"Muestra: {muestra}.\nCuándo usarla: {uso}."
                    )
                    n += 1
            destino = EXP / "presentaciones" / f"ponencia_{variante}_{tema}.pptx"
            prs.save(destino)
            print(destino.name, n, "videos", f"{destino.stat().st_size / 1e6:.0f} MB")


def stickers(por_slide=6, cols=3, ancho_px=1800):
    (EXP / "presentaciones").mkdir(parents=True, exist_ok=True)
    celda_w, celda_h = W / cols, H / 2
    for tema in ("oscuro", "claro"):
        prs = nueva(tema)
        for bloque, carpeta, piezas in BLOQUES:
            for k in range(0, len(piezas), por_slide):
                s = diapositiva(prs, tema)
                notas = [bloque]
                for j, (clase, muestra, _, _) in enumerate(piezas[k:k + por_slide]):
                    src = EXP / "png" / "stickers" / tema / carpeta / f"{clase}.png"
                    im = Image.open(src).convert("RGBA")
                    im.thumbnail((ancho_px, ancho_px), Image.LANCZOS)  # ligera para el pptx; el 8K queda en exports/
                    buf = io.BytesIO()
                    im.save(buf, "PNG", compress_level=6)
                    buf.seek(0)
                    esc = min((celda_w * 0.88) / im.width, (celda_h * 0.88) / im.height)
                    w, h = int(im.width * esc), int(im.height * esc)
                    x = int((j % cols) * celda_w + (celda_w - w) / 2)
                    y = int((j // cols) * celda_h + (celda_h - h) / 2)
                    pic = s.shapes.add_picture(buf, x, y, w, h)
                    pic._element.nvPicPr.cNvPr.set("descr", f"{nombre_legible(clase)}: {muestra}")
                    pic.name = clase
                    notas.append(f"{clase}: {muestra}")
                s.notes_slide.notes_text_frame.text = "\n".join(notas)
        destino = EXP / "presentaciones" / f"stickers_{tema}.pptx"
        prs.save(destino)
        print(destino.name, f"{destino.stat().st_size / 1e6:.0f} MB")


def montajes():
    (EXP / "montajes").mkdir(parents=True, exist_ok=True)
    for tema in ("oscuro", "claro"):
        listas = [(f"{i:02d}_{slug(bloque)}", [x for x in piezas], carpeta) for i, (bloque, carpeta, piezas) in enumerate(BLOQUES, 1)]
        listas.append(("reel_recomendadas", None, None))
        for nombre, piezas, carpeta in listas:
            if piezas is None:
                rutas = [EXP / c / tema / f"{p[0]}.mp4" for _, c, ps in BLOQUES for p in ps if p[3]]
            else:
                rutas = [EXP / carpeta / tema / f"{p[0]}.mp4" for p in piezas]
            rutas = [r for r in rutas if r.exists()]
            lista = EXP / "montajes" / f".lista_{nombre}_{tema}.txt"
            lista.write_text("".join(f"file '{r}'\n" for r in rutas))
            destino = EXP / "montajes" / f"{nombre}_{tema}.mp4"
            subprocess.run(["ffmpeg", "-loglevel", "error", "-y", "-f", "concat", "-safe", "0", "-i", str(lista),
                            "-c", "copy", str(destino)], check=True)
            lista.unlink()
            print(destino.name, len(rutas), "piezas")


if __name__ == "__main__":
    que = sys.argv[1] if len(sys.argv) > 1 else "todo"
    if que in ("presentaciones", "todo"):
        presentaciones()
    if que in ("stickers", "todo"):
        stickers()
    if que in ("montajes", "todo"):
        montajes()
