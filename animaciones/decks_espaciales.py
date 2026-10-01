"""Presentaciones «espaciales»: la misma secuencia y los mismos videos, en cualquier tema, en español o inglés.

No re-renderiza Manim: reutiliza los mp4 `_final` (terminan en el último cuadro, nunca en blanco) y, cuando el
tema lo pide, les recolorea el fondo con una LUT 3D de ffmpeg para que se fundan con la diapositiva
(exports/<archivo>/esp_<tema>[_en]/). Los textos dentro de los videos en inglés salen de renders propios
(IDIOMA=en ./render_todo.sh; carpetas oscuro_en / claro_en, ver traducciones_en.py).

TEMAS (temas_espaciales.py): orbita · mision · nebulosa · estacion, más los que se agreguen como plugin o JSON
en animaciones/temas/. Cada tema define fondos, tipografías, paleta y cómo se enmarca el video. Se puede
  · construir cualquier deck en cualquier tema,
  · mezclar temas en un deck (`--mezcla orbita,nebulosa`: uno por sección; `--mezcla-por diapositiva` alterna cada una),
  · crear temas nuevos (`python3 temas_espaciales.py --nuevo aurora --base orbita --tinte 140`).
Las diapositivas de solo texto (texto, cita, respaldo) llevan un sticker tomado del último cuadro de una pieza
relacionada; las de sección muestran miniaturas de los videos que vienen.

Uso:
  python3 decks_espaciales.py                              # tema oficial (Órbita) × 3 presentaciones en español
  python3 decks_espaciales.py --todos                      # todos los temas registrados
  python3 decks_espaciales.py orbita seminario             # filtra por tema y/o presentación
  python3 decks_espaciales.py --en                         # versión en inglés (las presentaciones que la tienen)
  python3 decks_espaciales.py --mezcla orbita,nebulosa seminario
  python3 decks_espaciales.py --solo-videos                # solo prepara videos recoloreados y pósters
  python3 decks_espaciales.py --lista                      # temas y presentaciones disponibles
Salida: exports/presentaciones/espaciales/<tema>/<archivo>_<tema>.pptx  (mezcla: espaciales/mezcla/<archivo>_<a>-<b>.pptx)
Fuentes: exports/presentaciones/espaciales/fuentes/ (instalar antes de abrir en PowerPoint).
"""
import io
import shutil
import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import numpy as np
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

sys.path.insert(0, str(Path(__file__).parent))
import catalogo  # noqa: E402
import catalogo_tesis  # noqa: E402
import fondos_espaciales as F  # noqa: E402
import temas_espaciales as TE  # noqa: E402
import traducciones_en as EN  # noqa: E402
from empaquetar_ponencia import EXP, nombre_legible  # noqa: E402
from tesis_decks import guion_de, palabras, video_final  # noqa: E402

SALIDA = EXP / "presentaciones" / "espaciales"
SW, SH = 13.333, 7.5
CARPETA = {p[0]: c for b in (catalogo.BLOQUES, catalogo_tesis.BLOQUES) for _, c, ps in b for p in ps}
MUESTRA = {p[0]: p[1] for b in (catalogo.BLOQUES, catalogo_tesis.BLOQUES) for _, _, ps in b for p in ps}
BLOQUE = {p[0]: t for b in (catalogo.BLOQUES, catalogo_tesis.BLOQUES) for t, _, ps in b for p in ps}

ZONA = TE.ZONA  # rectángulo del video según cómo se enmarca (x, y, w, h en pulgadas)

# ------------------------------------------------------------------ stickers de las diapositivas de solo texto
STICKERS = {
    "seminario": {
        "La idea de hoy, en una frase": "SateliteDeviceAgent",
        "Cómo citarlo con honestidad": "AstreaAconseja",
        "Nadie supervisa dieciséis mil satélites": "AritmeticaDeLaEscala",
        "Lenguaje normativo, en el texto original": "EstandaresEscribenAhora",
        "El mismo ejemplo, con números": "TermometroRoto",
        "Tres preguntas para leer cualquier comparación": "SignoQueInvierte",
        "Si el entorno no premia adaptarse, comparar algoritmos sobre él no significa nada.": "LazoAbiertoGana",
        "Lo mío es el uso, no la idea": "ProtocoloCompuertas",
        "Una tesis que empieza midiendo su propio instrumento": "ParMargenAdaptativo",
        "Para llevarse": "Tierra3DSatelites",
        "No estoy tratando de demostrar que la inteligencia artificial gana.": "CierreMedicionSignifica",
        "¿Los satélites ya piensan solos?": "AgenteAprende",
        "¿Y si el agente se equivoca allá arriba?": "NormaSinIA",
        "¿Esto sirve para algo en México?": "CoberturaHexagonal",
        "¿Por qué no un modelo de lenguaje para todo?": "ModeloArribaRLAbajo",
        "¿Cuánto te falta?": "CronogramaDoctoral",
    },
    "comite": {
        "Si el entorno no premia adaptarse, comparar algoritmos sobre él no significa nada.": "TermometroRoto",
        "¿Cómo gobernar de forma autónoma y resiliente": "DecisionConjunta",
        "Objetivo general y OE1′–OE5′: lo registrado": "PadaCuatroCajas",
        "P1 · ¿Y si los satélites son de operadores distintos?": "SupuestoA0",
        "P2 · ¿Su teorema es nuevo?": "DecPomdpUnificado",
        "P5 · ¿Por qué no está MAPPO en su comparación?": "G2bTresSemillas",
        "P6 · ¿Por qué QMIX y no un agente basado en modelos de lenguaje?": "ModeloArribaRLAbajo",
        "P11 · ¿Su compuerta no es vacua, si la holgura es no acotada?": "ParMargenAdaptativo",
        "P12 · ¿Cuántos entornos construyó hasta que uno pasó?": "RecorridoDelMargen",
        "P13 · Si el mixer no hace falta, ¿qué queda de SH3?": "VarianteSimpleNoEmpeora",
    },
    "divulgacion": {"Ruta de la charla": "TransferenciaHohmann"},
}


def _sticker_para(deck, clave):
    for k, v in STICKERS[deck].items():
        if clave == k or clave.startswith(k):
            return v
    return None


# ------------------------------------------------------------------ presentaciones → lista normalizada
# presentación → idioma → módulo con CFG y DIAPOS (la divulgación se arma desde ponencia_divulgacion.py)
DECKS = {"seminario": {"es": "tesis_seminario", "en": "tesis_seminario_en"},
         "comite": {"es": "tesis_comite"}, "divulgacion": {"es": "ponencia_divulgacion"}}


def _cargar(deck, idioma):
    """(cfg, [diapositivas como dict]); el sticker de las diapositivas de solo texto solo se asigna en español."""
    if deck == "divulgacion":
        import ponencia_divulgacion as P
        cfg = dict(archivo="redes_orbitales_divulgacion", kicker="PONENCIA DE DIVULGACIÓN  ·  DOS VOCES",
                   titulo=P.TITULO, lema=P.LEMA, autor="Alan Rosas Palacios  ·  Yuritzi Elena Ordaz Huerta",
                   instit="Co.De Aerospace  ·  codeaerospace.com", pie="Redes Orbitales · Co.De Aerospace",
                   cierre_titulo="El siguiente salto empieza en órbita", ritmo=P.RITMO, marca=True, tam_titulo=66)
        out = []
        for d in P.DIAPOS:
            t, q, g = d[0], d[1], d[-1]
            x = dict(tipo=t, quien=q, guion=g, cuidado="")
            if t == "texto":
                x.update(titulo=d[2], sub=d[3], puntos=d[4], sticker=d[5] or _sticker_para(deck, d[2]))
            elif t == "seccion":
                x.update(titulo=d[2], sub=d[3])
            elif t == "video":
                x.update(clase=d[2], titulo=d[3], sub=d[4])
            out.append(x)
        return cfg, out
    mod = __import__(DECKS[deck][idioma])
    cfg = dict(mod.CFG)
    cfg["marca"] = False
    out = []
    for d in mod.DIAPOS:
        g, c = guion_de(d)
        t = d[0]
        x = dict(tipo=t, guion=g, cuidado=c)
        if t == "seccion":
            x.update(titulo=d[1], sub=d[2])
        elif t == "video":
            x.update(clase=d[1], titulo=d[2], sub=d[3])
        elif t in ("texto", "respaldo"):
            x.update(titulo=d[1], sub=d[2], puntos=d[3], sticker=_sticker_para(deck, d[1]) if idioma == "es" else None)
        elif t == "cita":
            x.update(frase=d[1], sub=d[2], sticker=_sticker_para(deck, d[1]) if idioma == "es" else None)
        out.append(x)
    return cfg, out


def presentacion(deck, idioma="es"):
    """(cfg, [diapositivas como dict]) en el idioma pedido. La versión en inglés hereda los stickers de la española
    (misma secuencia, misma posición) y la advertencia `idioma` para el constructor."""
    if idioma not in DECKS[deck]:
        raise KeyError(f"{deck} no tiene versión en {idioma!r} (tiene: {', '.join(DECKS[deck])})")
    cfg, out = _cargar(deck, idioma)
    if idioma != "es":
        _, es = _cargar(deck, "es")
        assert len(es) == len(out) and all(a["tipo"] == b["tipo"] for a, b in zip(es, out)), "la traducción no calza con el original"
        for a, b in zip(es, out):
            if "sticker" in b:
                b["sticker"] = a.get("sticker")
    cfg["idioma"] = idioma
    return cfg, out


# ------------------------------------------------------------------ videos adaptados al entorno
def _ffmpeg(*a):
    subprocess.run(["ffmpeg", "-v", "error", "-y", *map(str, a)], check=True)


def _ultimo_cuadro(mp4, matriz):
    raw = subprocess.run(["ffmpeg", "-v", "error", "-sseof", "-0.1", "-i", str(mp4), "-frames:v", "1", "-vf",
                          f"scale=in_color_matrix={matriz}:in_range=tv,format=rgb24", "-f", "rawvideo", "-"],
                         capture_output=True, check=True).stdout
    return np.frombuffer(raw, np.uint8).reshape(-1, 1920, 3)[-1080:]


def _lut(viejo, nuevo, ruta):
    """LUT 3D: los colores cercanos al fondo viejo se desplazan hacia el nuevo; el resto queda igual."""
    g = np.linspace(0, 255, 33)
    b, gg, r = np.meshgrid(g, g, g, indexing="ij")
    c = np.stack([r, gg, b], -1).reshape(-1, 3)
    d = np.linalg.norm(c - viejo, axis=1)
    t = np.clip((d - 6) / (60 - 6), 0, 1)
    w = 1 - t * t * (3 - 2 * t)
    o = np.clip(c + (nuevo - viejo)[None] * w[:, None], 0, 255) / 255
    ruta.write_text("LUT_3D_SIZE 33\n" + "\n".join(f"{x:.6f} {y:.6f} {z:.6f}" for x, y, z in o))


def dir_modo(tema, clase, idioma):
    """Carpeta de videos de origen y si es la versión en inglés.
    Si el tema tiene render nativo de la pieza (Manim con su fuente y paleta: t_<tema>[_en], ver render_todo.sh TEMA_VISUAL) se usa;
    si no, el video oscuro|claro con el fondo recoloreado. Las piezas sin texto con letras sirven tal cual en inglés."""
    en = idioma == "en" and clase not in EN.SIN_TEXTO
    nativo = f"t_{tema.id}" + ("_en" if en else "")
    if tema.manim and (EXP / CARPETA[clase] / nativo / f"{clase}.mp4").exists():
        return nativo, en
    return tema.modo + ("_en" if en else ""), en


def video_para(tema, clase, idioma="es"):
    """(mp4, jpg del póster) del video listo para el tema (y el idioma)."""
    if isinstance(tema, str):
        tema = TE.obtener(tema)
    carpeta = CARPETA[clase]
    tdir, en = dir_modo(tema, clase, idioma)
    src = video_final(tdir, clase, carpeta)
    sub = f"esp_{tema.id}" + ("_en" if en else "")
    if tema.pantalla is None:
        dst, matriz = src, "bt601"
    else:
        dst = EXP / carpeta / sub / f"{clase}.mp4"
        matriz = "bt709"
        if not (dst.exists() and dst.stat().st_mtime >= src.stat().st_mtime):
            dst.parent.mkdir(parents=True, exist_ok=True)
            cuadro = _ultimo_cuadro(src, "bt601").reshape(-1, 3)
            v, n = np.unique(cuadro, axis=0, return_counts=True)
            viejo = v[np.argmax(n)].astype(float)
            nuevo = F.rgb(tema.pantalla) * 255
            lut = dst.with_suffix(".cube")
            _lut(viejo, nuevo, lut)
            _ffmpeg("-i", src, "-vf", f"scale=in_color_matrix=bt601:in_range=tv,format=rgb24,lut3d=file={lut}:interp=tetrahedral,"
                    "scale=out_color_matrix=bt709:out_range=tv,format=yuv420p", "-an", "-c:v", "libx264", "-crf", "17",
                    "-preset", "fast", "-colorspace", "bt709", "-color_primaries", "bt709", "-color_trc", "bt709",
                    "-color_range", "tv", "-movflags", "+faststart", dst.with_suffix(".tmp.mp4"))
            dst.with_suffix(".tmp.mp4").replace(dst)
            lut.unlink()
    jpg = EXP / carpeta / sub / f"{clase}.poster.jpg"
    if not (jpg.exists() and jpg.stat().st_mtime >= dst.stat().st_mtime):
        jpg.parent.mkdir(parents=True, exist_ok=True)
        Image.fromarray(_ultimo_cuadro(dst, matriz)).save(jpg, "JPEG", quality=90)
    return dst, jpg


def sticker_final(tema, clase):
    """Sticker con el ESTADO FINAL de la pieza: último cuadro legible del video `_final`, fondo vuelto transparente.

    Los PNG de exports/png/stickers se tomaron a media animación y algunos muestran cifras intermedias
    (p. ej. PlanosOrbitales 96 % / 24 satélites cuando el video termina en 99 % / 36); este muestra lo mismo que
    queda en pantalla al acabar el video. Mate de diferencia: alfa según la distancia al color de fondo y
    des-premultiplicado para recuperar el color original de bordes y rellenos translúcidos.
    """
    carpeta = CARPETA[clase]
    dst = EXP / "png" / "stickers_finales" / tema / carpeta / f"{clase}.png"
    src = video_final(tema, clase, carpeta)
    if dst.exists() and dst.stat().st_mtime >= src.stat().st_mtime:
        return dst
    c = _ultimo_cuadro(src, "bt601").astype(np.float32)
    v, n = np.unique(c.reshape(-1, 3).astype(np.uint8), axis=0, return_counts=True)
    fondo = v[np.argmax(n)].astype(np.float32)
    dif = np.abs(c - fondo).max(-1)
    a = np.clip((dif - 7) / 60, 0, 1)
    col = fondo + (c - fondo) / np.maximum(a, 1e-3)[..., None]
    rgba = np.dstack([np.clip(col, 0, 255), a * 255]).astype(np.uint8)
    im = Image.fromarray(rgba, "RGBA")
    bb = im.getchannel("A").point(lambda x: 255 if x > 40 else 0).getbbox()
    if bb:
        m = 12
        im = im.crop((max(0, bb[0] - m), max(0, bb[1] - m), min(1920, bb[2] + m), min(1080, bb[3] + m)))
    dst.parent.mkdir(parents=True, exist_ok=True)
    im.save(dst, "PNG", optimize=True)
    return dst


def duracion(mp4):
    return float(subprocess.check_output(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0",
                                          str(mp4)]).decode())


# ------------------------------------------------------------------ utilidades de pptx
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def I(v):  # noqa: E743  pulgadas → EMU
    return Emu(int(v * 914400))


def color(h):
    return RGBColor.from_string(h.lstrip("#").upper())


def _alfa(srgb, a):
    for x in srgb.findall(qn("a:alpha")):
        srgb.remove(x)
    etree.SubElement(srgb, qn("a:alpha")).set("val", str(int(a * 100000)))


def poner_fondo(s, ruta):
    _, rid = s.part.get_or_add_image_part(str(ruta))
    csld = s._element.find(qn("p:cSld"))
    viejo = csld.find(qn("p:bg"))
    if viejo is not None:
        csld.remove(viejo)
    bg = etree.fromstring(
        f'<p:bg xmlns:p="{NS_P}" xmlns:a="{NS_A}" xmlns:r="{NS_R}"><p:bgPr><a:blipFill dpi="0" rotWithShape="1">'
        f'<a:blip r:embed="{rid}"/><a:srcRect/><a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>')
    csld.insert(0, bg)


def caja(s, x, y, w, h, relleno=None, a_relleno=1.0, borde=None, a_borde=1.0, grosor=1.0, forma=MSO_SHAPE.RECTANGLE,
         radio=None, sombra=None):
    sh = s.shapes.add_shape(forma, I(x), I(y), I(w), I(h))
    sh.shadow.inherit = False
    if relleno:
        sh.fill.solid()
        sh.fill.fore_color.rgb = color(relleno)
        if a_relleno < 1:
            _alfa(sh._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr")), a_relleno)
    else:
        sh.fill.background()
    if borde:
        sh.line.color.rgb = color(borde)
        sh.line.width = Pt(grosor)
        if a_borde < 1:
            _alfa(sh._element.spPr.find(qn("a:ln")).find(qn("a:solidFill")).find(qn("a:srgbClr")), a_borde)
    else:
        sh.line.fill.background()
    if radio is not None:
        sh.adjustments[0] = radio
    if sombra:
        spPr = sh._element.spPr
        for e in spPr.findall(qn("a:effectLst")):
            spPr.remove(e)
        c, a, blur, dist = sombra
        spPr.append(etree.fromstring(
            f'<a:effectLst xmlns:a="{NS_A}"><a:outerShdw blurRad="{int(blur * 12700)}" dist="{int(dist * 12700)}" dir="5400000" '
            f'algn="t" rotWithShape="0"><a:srgbClr val="{c.lstrip("#")}"><a:alpha val="{int(a * 100000)}"/></a:srgbClr>'
            f'</a:outerShdw></a:effectLst>'))
    return sh


def degradado(s, x, y, w, h, c, a0=1.0, a1=0.0, ang=0, radial=False):
    """Rectángulo con degradado lineal (ang en grados) o radial de c con alfa a0 → a1."""
    sh = caja(s, x, y, w, h, relleno=c, forma=MSO_SHAPE.OVAL if radial else MSO_SHAPE.RECTANGLE)
    spPr = sh._element.spPr
    sf = spPr.find(qn("a:solidFill"))
    i = list(spPr).index(sf)
    spPr.remove(sf)
    c = c.lstrip("#")
    forma = ('<a:path path="circle"><a:fillToRect l="50000" t="50000" r="50000" b="50000"/></a:path>' if radial
             else f'<a:lin ang="{int(ang * 60000)}" scaled="0"/>')
    spPr.insert(i, etree.fromstring(
        f'<a:gradFill xmlns:a="{NS_A}" rotWithShape="1"><a:gsLst>'
        f'<a:gs pos="0"><a:srgbClr val="{c}"><a:alpha val="{int(a0 * 100000)}"/></a:srgbClr></a:gs>'
        f'<a:gs pos="100000"><a:srgbClr val="{c}"><a:alpha val="{int(a1 * 100000)}"/></a:srgbClr></a:gs>'
        f'</a:gsLst>{forma}</a:gradFill>'))
    return sh


def texto(s, x, y, w, h, contenido, tam, col, fuente, negrita=False, alinear="l", spc=None, ancla="t", interlinea=None,
          alfa=None):
    """Caja de texto sin márgenes. `contenido` = str o lista de (texto, fuente, tam, color, negrita[, spc])."""
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[ancla]
    p = tf.paragraphs[0]
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[alinear]
    if interlinea:
        p.line_spacing = interlinea
    runs = contenido if isinstance(contenido, list) else [(contenido, fuente, tam, col, negrita, spc)]
    for rr in runs:
        t, f, z, c, b = rr[:5]
        sp = rr[5] if len(rr) > 5 else None
        r = p.add_run()
        r.text = t
        r.font.name, r.font.size, r.font.bold = f, Pt(z), b
        r.font.color.rgb = color(c)
        if sp:
            r.font._rPr.set("spc", str(int(sp * 100)))
        if alfa is not None:
            _alfa(r.font._rPr.find(qn("a:solidFill")).find(qn("a:srgbClr")), alfa)
    return tb


def lineas_est(t, tam, ancho, k_car):
    """Líneas estimadas de un texto de `tam` pt en `ancho` pulgadas (k_car = ancho medio del carácter en em)."""
    por_linea = max(8, int(ancho * 72 / (tam * k_car)))
    n, actual = 1, 0
    for pal in t.split():
        if actual + len(pal) + (1 if actual else 0) > por_linea:
            n, actual = n + 1, len(pal)
        else:
            actual += len(pal) + (1 if actual else 0)
    return n


def ajustar(t, ancho, alto, tams, k_car, inter=1.12):
    """Primer tamaño de `tams` con el que el texto cabe en alto (pulgadas)."""
    for z in tams:
        if lineas_est(t, z, ancho, k_car) * z * inter / 72 <= alto:
            return z
    return tams[-1]


# ------------------------------------------------------------------ piezas de diseño
class Diseno:
    """Construye un deck. `temas`: lista de ids o Tema; con más de uno se mezclan (un tema por grupo, ver `reparto`):
    `seccion` = un tema por sección (en el comité, que no tiene secciones, por bloque de piezas), `diapositiva` = alterna en cada una."""

    def __init__(self, temas, deck, cfg, diapos, idioma="es", reparto="seccion"):
        temas = [temas] if isinstance(temas, (str, TE.Tema)) else list(temas)
        self.temas = [TE.obtener(t) if isinstance(t, str) else t for t in temas]
        self.deck, self.cfg, self.diapos, self.idioma, self.reparto = deck, cfg, diapos, idioma, reparto
        self.usar(self.temas[0])
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = I(SW), I(SH)
        self._stk = {}
        self._contexto()

    def usar(self, tema):
        self.E, self.ent = tema, tema.id
        self.C, self.Fu = tema.color, tema.fuentes

    # ---- idioma
    def t(self, clave):
        return EN.UI.get(clave, clave) if self.idioma == "en" else clave

    def nombre(self, clase):
        return EN.NOMBRES.get(clase, nombre_legible(clase)) if self.idioma == "en" else nombre_legible(clase)

    def bloque(self, b):
        return EN.BLOQUES.get(b, b) if self.idioma == "en" and b else b

    def tema_de(self, d, n):
        if len(self.temas) == 1:
            return self.temas[0]
        return self.temas[(n - 1 if self.reparto == "diapositiva" else d["_grupo"]) % len(self.temas)]

    # contexto: sección, bloque, hora T+, videos por sección
    def _contexto(self):
        ritmo = self.cfg["ritmo"]
        t, sec, tit_sec = 0.0, 0, None
        self.nsec = sum(1 for d in self.diapos if d["tipo"] == "seccion")
        self.videos_de = {}
        ultimo_bloque = None
        for i, d in enumerate(self.diapos):
            if d["tipo"] == "seccion":
                sec += 1
                tit_sec = d["titulo"]
                self.videos_de[sec] = []
            if d["tipo"] == "video":
                ultimo_bloque = BLOQUE[d["clase"]]
                if sec:
                    self.videos_de[sec].append(d["clase"])
            d["_sec"], d["_tsec"], d["_t"] = sec, tit_sec, t
            if d["tipo"] != "respaldo":
                t += palabras(d["guion"]) / ritmo
            d["_bloque"] = ultimo_bloque
        # sin secciones (comité): el bloque de la pieza siguiente da el contexto
        siguiente = None
        for d in reversed(self.diapos):
            if d["tipo"] == "video":
                siguiente = BLOQUE[d["clase"]]
            d["_bloque_sig"] = siguiente or d["_bloque"]
        orden = []
        for d in self.diapos:
            b = d["_bloque_sig"]
            if b not in orden:
                orden.append(b)
        for d in self.diapos:  # grupo de mezcla: la sección (las diapositivas previas van con la 1.ª) o, sin secciones, el bloque
            d["_grupo"] = max(d["_sec"], 1) - 1 if self.nsec else orden.index(d["_bloque_sig"])

    def kicker_de(self, d):
        if d["tipo"] == "respaldo":
            return self.t("RESPALDO  ·  PREGUNTAS")
        if d["_sec"]:
            return f"{d['_sec']:02d}  ▸  {d['_tsec'].upper()}"
        b = self.bloque(d["_bloque"] if d["tipo"] == "video" else d["_bloque_sig"])
        return (b or self.cfg["kicker"]).upper()

    def fondo(self, tipo, var=0):
        return self.E.fondo(tipo, var)

    # ---------------------------------------------------------- comunes
    def nueva(self, tipo, var=0):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        poner_fondo(s, self.fondo(tipo, var))
        return s

    def kicker(self, s, x, y, w, t, col=None, tam=10.5):
        texto(s, x, y, w, 0.3, t, tam, col or self.C["acento"], self.Fu["etiqueta"], spc=2.4)

    def pie(self, s, n, d):
        C, et = self.C, self.Fu["etiqueta"]
        izq = self.cfg["pie"].upper()
        if d.get("quien"):
            izq = f"{self.t('VOZ')} ▸ {d['quien'].upper()}     ·     " + izq
        texto(s, 0.7, 7.1, 8.5, 0.25, izq, 8, C["tenue"], et, spc=1.5)
        total = len(self.diapos)
        t = d["_t"]
        hora = self.t("RESPALDO") if d["tipo"] == "respaldo" else f"T+{int(t):02d}:{int((t % 1) * 60):02d}"
        texto(s, 9.0, 7.1, 3.63, 0.25, [(hora, et, 8, C["acento"], False, 1.5), (f"     {n:02d} / {total:02d}", et, 8, C["tenue"], False, 1.5)],
              8, C["tenue"], et, alinear="r")
        if self.E.esquinas:
            for (x, y, sx, sy) in ((0.2, 0.2, 1, 1), (SW - 0.2, 0.2, -1, 1), (0.2, SH - 0.2, 1, -1), (SW - 0.2, SH - 0.2, -1, -1)):
                self.escuadra(s, x, y, sx, sy, 0.24, C["acento2"], 0.8)

    def escuadra(self, s, x, y, sx, sy, L, col, a=1.0, g=0.025):
        caja(s, x if sx > 0 else x - L, y - (g if sy < 0 else 0), L, g, relleno=col, a_relleno=a)
        caja(s, x - (g if sx < 0 else 0), y if sy > 0 else y - L, g, L, relleno=col, a_relleno=a)

    def notas(self, s, d, extra=""):
        g = f"[{d['quien']}] " + d["guion"] if d.get("quien") else d["guion"]
        if d.get("cuidado"):
            g += f"\n\n⚠️ {self.t('CUIDADO')}: {d['cuidado']}"
        s.notes_slide.notes_text_frame.text = g + extra

    def titulo(self, s, x, y, w, t, tams, alto, alinear="l", col=None):
        k = self.E["k_titulo"]
        z = ajustar(t, w, alto, [round(v * k) for v in tams], self.E["ancho_car"]["titulo"], 1.0)
        texto(s, x, y, w, alto, t, z, col or self.C["tinta"], self.Fu["titulo"], negrita=True, alinear=alinear,
              interlinea=0.95)
        return z, lineas_est(t, z, w, self.E["ancho_car"]["titulo"])

    def regla(self, s, x, y, w):
        degradado(s, x, y, w, 0.018, self.C["linea"], 0.9, 0.0, 0)

    def sticker(self, s, clase, x, y, w, h, pie=True):
        tdir, _ = dir_modo(self.E, clase, self.idioma)
        if (tdir, clase) not in self._stk:
            im = Image.open(sticker_final(tdir, clase))
            b = io.BytesIO()
            im.save(b, "PNG", compress_level=6)
            self._stk[(tdir, clase)] = (b.getvalue(), im.size)
        data, (iw, ih) = self._stk[(tdir, clase)]
        alto_pie = 0.35 if pie else 0
        esc = min(w / iw, (h - alto_pie) / ih)
        pw, ph = iw * esc, ih * esc
        px, py = x + (w - pw) / 2, y + (h - alto_pie - ph) / 2
        if self.E.esquinas:
            for (cx, cy, sx, sy) in ((px - 0.12, py - 0.12, 1, 1), (px + pw + 0.12, py - 0.12, -1, 1),
                                     (px - 0.12, py + ph + 0.12, 1, -1), (px + pw + 0.12, py + ph + 0.12, -1, -1)):
                self.escuadra(s, cx, cy, sx, sy, 0.22, self.C["acento"], 0.9)
        pic = s.shapes.add_picture(io.BytesIO(data), I(px), I(py), I(pw), I(ph))
        pic._element.nvPicPr.cNvPr.set("descr", self.nombre(clase) + ("" if self.idioma == "en" else f": {MUESTRA.get(clase, '')}"))
        if pie:
            texto(s, x, py + ph + 0.14, w, 0.25, f"FIG  ·  {self.nombre(clase).upper()}", 7.5, self.C["tenue"],
                  self.Fu["etiqueta"], alinear="c", spc=2)

    def vinetas(self, s, puntos, x, y, w, h):
        k = self.E["ancho_car"]["cuerpo"]
        sangria = 0.62
        for z in (22, 20, 19, 18, 17, 16, 15, 14):
            lin = sum(lineas_est(t, z, w - sangria, k) for t in puntos)
            if lin * z * 1.2 / 72 + (len(puntos) - 1) * z * 0.8 / 72 <= h:
                break
        tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, t in enumerate(puntos):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(I(sangria))))
            pPr.set("indent", str(-int(I(sangria))))
            p.space_after = Pt(z * 0.8)
            p.line_spacing = 1.08
            r1 = p.add_run()
            r1.text = f"{i + 1:02d}\t"
            r1.font.name, r1.font.size, r1.font.bold = self.Fu["etiqueta"], Pt(z * 0.7), True
            r1.font.color.rgb = color(self.C["acento"])
            r2 = p.add_run()
            r2.text = t
            r2.font.name, r2.font.size = self.Fu["cuerpo"], Pt(z)
            r2.font.color.rgb = color(self.C["tinta"])
        return z

    # ---------------------------------------------------------- tipos de diapositiva
    def portada(self, d, n):
        s = self.nueva("portada")
        C, Fu, cfg = self.C, self.Fu, self.cfg
        w = self.E["portada_ancho"]
        x = 0.85
        if cfg.get("marca"):
            texto(s, x, 0.55, 6, 0.35, "CO.DE AEROSPACE", 13, C["tinta"], "Orbitron", negrita=True, spc=5)
        self.kicker(s, x, 1.25, w, cfg["kicker"], tam=11)
        tam = cfg.get("tam_titulo", 54) + 6
        z, nl = self.titulo(s, x, 1.62, w, cfg["titulo"], [tam, tam - 6, tam - 12, 40, 36], 2.3)
        y = 1.62 + min(nl, 3) * z * 1.12 / 72 + 0.2
        texto(s, x, y, w, 0.9, cfg["lema"], 24, C["acento"], Fu["cuerpo"])
        y += 0.95
        self.regla(s, x, y, min(w, 6.5))
        texto(s, x, y + 0.22, w, 0.45, cfg["autor"], 20, C["tinta"], Fu["titulo"], negrita=True)
        texto(s, x, y + 0.66, w, 0.6, cfg["instit"], 13, C["tenue"], Fu["cuerpo"])
        if self.E.esquinas:
            self.pie_portada_mision(s)
        self.notas(s, d)
        return s

    def pie_portada_mision(self, s):
        C = self.C
        texto(s, 0.85, 6.72, 7, 0.25, self.t("LAT_LON") if self.idioma == "en" else "LAT 19.43° N   LON 99.13° O   ·   CIUDAD DE MÉXICO", 8, C["acento2"], self.Fu["etiqueta"], spc=2)
        for (x, y, sx, sy) in ((0.2, 0.2, 1, 1), (SW - 0.2, 0.2, -1, 1), (0.2, SH - 0.2, 1, -1), (SW - 0.2, SH - 0.2, -1, -1)):
            self.escuadra(s, x, y, sx, sy, 0.24, C["acento2"], 0.8)

    def cierre(self, d, n):
        s = self.nueva("cierre")
        C, Fu, cfg = self.C, self.Fu, self.cfg
        derecha = self.E.cierre_derecha
        x, w = (6.2, 6.4) if derecha else (0.85, 9.5)
        self.kicker(s, x, 1.55, w, self.t("GRACIAS  ·  PREGUNTAS"), tam=11)
        z, nl = self.titulo(s, x, 1.95, w, cfg["cierre_titulo"], [52, 46, 40, 36], 2.2)
        y = 1.95 + min(nl, 3) * z * 1.12 / 72 + 0.3
        self.regla(s, x, y, min(w, 6))
        texto(s, x, y + 0.25, w, 0.45, cfg["autor"], 18, C["tinta"], Fu["titulo"], negrita=True)
        texto(s, x, y + 0.68, w, 0.6, cfg["instit"], 13, C["tenue"], Fu["cuerpo"])
        if self.E.esquinas:
            self.pie_portada_mision(s)
        self.notas(s, d)
        return s

    def seccion(self, d, n):
        sec = d["_sec"]
        s = self.nueva("seccion", sec - 1)
        C, Fu = self.C, self.Fu
        x, w = 0.85, 7.3
        texto(s, x - 0.05, 0.6, 5, 2.0, f"{sec:02d}", round(120 * self.E["k_titulo"]), C["acento"], Fu["titulo"],
              negrita=True, alfa=0.9)
        self.kicker(s, x, 2.75, w, f"{self.t('SECCIÓN')} {sec:02d} / {self.nsec:02d}")
        z, nl = self.titulo(s, x, 3.1, w, d["titulo"], [48, 42, 38, 34], 1.9)
        y = 3.1 + min(nl, 3) * z * 1.12 / 72 + 0.12
        texto(s, x, y, w, 0.9, d["sub"], 20, C["acento2"], Fu["cuerpo"])
        # miniaturas de lo que viene
        vids = self.videos_de.get(sec, [])[:6]
        if vids:
            texto(s, x, 5.55, 6, 0.25, f"{len(self.videos_de[sec])} {self.t('PIEZAS EN ESTA SECCIÓN')}", 8, C["tenue"], Fu["etiqueta"], spc=2)
            tw, th = 1.3, 1.3 * 9 / 16
            for i, c in enumerate(vids):
                _, jpg = video_para(self.E, c, self.idioma)
                xx = x + i * (tw + 0.14)
                caja(s, xx - 0.03, 5.87, tw + 0.06, th + 0.06, relleno=C["panel"], borde=C["linea"], a_borde=0.45, grosor=0.75)
                s.shapes.add_picture(str(jpg), I(xx), I(5.9), I(tw), I(th))
        self.pie(s, n, d)
        self.notas(s, d)
        return s

    def cabecera(self, s, d, x=0.7, w=11.93):
        self.kicker(s, x, 0.32, w, self.kicker_de(d))
        z, nl = self.titulo(s, x, 0.58, w, d["titulo"], [32, 29, 26, 23], 0.55)
        texto(s, x, 1.1, w, 0.35, d["sub"], 15.5, self.C["acento2"], self.Fu["cuerpo"])

    def video(self, d, n):
        s = self.nueva("video")
        C, Fu = self.C, self.Fu
        clase = d["clase"]
        mp4, jpg = video_para(self.E, clase, self.idioma)
        x, y, w, h = self.E.zona
        self.cabecera(s, d)
        modo = self.E.video
        if modo == "monitor":
            pad = 0.07
            caja(s, x - pad, y - pad, w + 2 * pad, h + 2 * pad, relleno=C["panel"], borde=C["acento2"], a_borde=0.35, grosor=0.75)
            for (cx, cy, sx, sy) in ((x - pad - 0.08, y - pad - 0.08, 1, 1), (x + w + pad + 0.08, y - pad - 0.08, -1, 1),
                                     (x - pad - 0.08, y + h + pad + 0.08, 1, -1), (x + w + pad + 0.08, y + h + pad + 0.08, -1, -1)):
                self.escuadra(s, cx, cy, sx, sy, 0.3, C["acento"], 1.0, 0.03)
            et = Fu["etiqueta"]
            texto(s, x - pad, y - pad - 0.27, 6, 0.22, f"{self.t('CANAL')} {n:02d}  ▸  {self.nombre(clase).upper()}", 8, C["acento2"], et, spc=1.5)
            texto(s, x + w - 4 + pad, y - pad - 0.27, 4, 0.22, [("●  ", et, 8, C["calido"], False), (f"{duracion(mp4):04.1f} s", et, 8, C["tenue"], False, 1.5)],
                  8, C["tenue"], et, alinear="r")
        elif modo == "tarjeta":
            pad = 0.13
            claro = self.E.modo == "claro"
            caja(s, x - pad, y - pad, w + 2 * pad, h + 2 * pad, relleno=C["panel"], borde=None if claro else C["linea"],
                 a_borde=0.45, grosor=1.0, forma=MSO_SHAPE.ROUNDED_RECTANGLE, radio=0.025,
                 sombra=("#1E3A8A", 0.18, 28, 6) if claro else ("#000000", 0.6, 40, 0))
        s.shapes.add_movie(str(mp4), I(x), I(y), I(w), I(h), poster_frame_image=str(jpg), mime_type="video/mp4")
        self.pie(s, n, d)
        if self.idioma == "en":
            nota = f"\n\n(Video: {self.nombre(clase)}. {EN.UI['video_nota']})"
        else:
            nota = (f"\n\n(Video: {nombre_legible(clase)} — {MUESTRA.get(clase, '')}. Clic para reproducir; "
                    "al terminar queda el último cuadro.)")
        self.notas(s, d, nota)
        return s

    def texto_(self, d, n):
        s = self.nueva("contenido")
        C = self.C
        x, w = 0.7, 6.35
        self.kicker(s, x, 0.55, 11.9, self.kicker_de(d))
        z, nl = self.titulo(s, x, 0.85, 11.9, d["titulo"], [38, 34, 30, 27], 1.2)
        y = 0.85 + min(nl, 2) * z * 1.12 / 72 + 0.06
        texto(s, x, y, 11.9, 0.4, d["sub"], 17, C["acento2"], self.Fu["cuerpo"])
        y += 0.62
        self.regla(s, x, y, 5.5)
        y += 0.3
        self.vinetas(s, d["puntos"], x, y, w, 6.85 - y)
        if d.get("sticker"):
            self.sticker(s, d["sticker"], 7.35, max(1.9, y - 0.2), 5.55, 6.9 - max(1.9, y - 0.2))
        self.pie(s, n, d)
        self.notas(s, d)
        return s

    def cita(self, d, n):
        s = self.nueva("contenido")
        C, Fu = self.C, self.Fu
        x, w = 0.95, 6.55
        self.kicker(s, 0.7, 0.55, 11.9, self.kicker_de(d))
        texto(s, 0.55, 0.75, 1.6, 1.6, "“", 150, C["acento"], Fu["display"] if Fu["display"] != "Orbitron" else Fu["titulo"],
              negrita=True, alfa=0.35)
        frase = d["frase"]
        k = self.E["k_titulo"]
        z = ajustar(frase, w, 3.6, [round(v * k) for v in (44, 40, 36, 32, 28, 25, 23)], self.E["ancho_car"]["titulo"], 1.05)
        texto(s, x, 1.75, w, 3.7, frase, z, C["tinta"], Fu["titulo"], negrita=True, ancla="m", interlinea=1.0)
        if d.get("sub"):
            self.regla(s, x, 5.6, 3.5)
            texto(s, x, 5.78, w, 0.9, d["sub"], 18, C["acento2"], Fu["cuerpo"])
        if d.get("sticker"):
            self.sticker(s, d["sticker"], 7.75, 1.3, 5.2, 5.4)
        self.pie(s, n, d)
        self.notas(s, d)
        return s

    def construir(self):
        for n, d in enumerate(self.diapos, 1):
            self.usar(self.tema_de(d, n))
            {"portada": self.portada, "cierre": self.cierre, "seccion": self.seccion, "video": self.video,
             "texto": self.texto_, "respaldo": self.texto_, "cita": self.cita}[d["tipo"]](d, n)
        if len(self.temas) == 1:
            destino = SALIDA / self.temas[0].id / f"{self.cfg['archivo']}_{self.temas[0].id}.pptx"
        else:
            destino = SALIDA / "mezcla" / f"{self.cfg['archivo']}_{'-'.join(t.id for t in self.temas)}.pptx"
        destino.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(destino)
        print(f"  {destino.relative_to(EXP)}  ·  {len(self.prs.slides)} diapositivas  ·  {destino.stat().st_size / 1e6:.0f} MB",
              flush=True)
        return destino


# ------------------------------------------------------------------ orquestación
def construir(temas, deck, idioma="es", reparto="seccion"):
    cfg, ds = presentacion(deck, idioma)
    return Diseno(temas, deck, cfg, ds, idioma, reparto).construir()


def preparar_videos(temas, decks, idioma="es"):
    """Recolorea los videos (y saca pósters) que necesitan los temas × presentaciones pedidos."""
    tareas = sorted({(t.id, d["clase"]) for t in temas for k in decks
                     for d in presentacion(k, idioma)[1] if d["tipo"] == "video"})
    print(f"videos a preparar: {len(tareas)}", flush=True)
    hechos = 0

    def uno(t):
        return video_para(TE.obtener(t[0]), t[1], idioma)

    with ThreadPoolExecutor(max_workers=3) as ex:
        for _ in ex.map(uno, tareas):
            hechos += 1
            if hechos % 20 == 0:
                print(f"  {hechos}/{len(tareas)}", flush=True)


def copiar_fuentes():
    dst = SALIDA / "fuentes"
    dst.mkdir(parents=True, exist_ok=True)
    src = Path.home() / ".local/share/fonts/codeaerospace"
    for f in sorted(src.glob("*.ttf")):
        shutil.copy2(f, dst / f.name)
    filas = "\n".join(f"| {t.nombre} | {t.fuentes['titulo']} | {t.fuentes['cuerpo']} | {t.fuentes['etiqueta']} |"
                      for t in TE.TEMAS.values())
    (dst / "LEEME.md").write_text(
        "# Fuentes de las presentaciones espaciales\n\n"
        "Instalar TODAS antes de abrir los .pptx (Windows: seleccionar los .ttf → clic derecho → «Instalar para todos los usuarios»; "
        "macOS: doble clic → Instalar). Si no están instaladas, PowerPoint sustituye por otra fuente y los títulos cambian de ancho.\n\n"
        "Para presentar en otra computadora sin instalarlas: PowerPoint (Windows) → Archivo → Opciones → Guardar → "
        "«Incrustar fuentes en el archivo».\n\n"
        "| Tema | Títulos | Texto | Etiquetas |\n|---|---|---|---|\n" + filas + "\n\n"
        "Todas son de Google Fonts con licencia SIL Open Font License (uso libre, incluso comercial). "
        "Rajdhani, DM Sans, Space Mono y Orbitron son las que usa codeaerospace.com. "
        "Un tema nuevo que use otra fuente: súmala aquí (copia el .ttf a ~/.local/share/fonts/codeaerospace/).\n", encoding="utf-8")


def main(argv):
    import argparse
    ap = argparse.ArgumentParser(description="Presentaciones espaciales: temas, idioma y mezcla.")
    ap.add_argument("nombres", nargs="*", help="ids de tema y/o de presentación (seminario, comite, divulgacion)")
    ap.add_argument("--en", action="store_true", help="versión en inglés (solo las presentaciones que la tienen)")
    ap.add_argument("--mezcla", help="temas separados por coma: se reparten en el deck (uno por sección)")
    ap.add_argument("--mezcla-por", choices=["seccion", "diapositiva"], default="seccion")
    ap.add_argument("--todos", action="store_true", help="todos los temas registrados (por defecto solo el oficial: Órbita)")
    ap.add_argument("--solo-videos", action="store_true")
    ap.add_argument("--lista", action="store_true")
    a = ap.parse_args(argv)
    idioma = "en" if a.en else "es"
    if a.lista:
        for t in TE.TEMAS.values():
            print(f"{t.id:10s} {t.nombre:18s} video={t.video:8s} {t.descripcion}")
        for k, v in DECKS.items():
            print(f"{k:12s} idiomas: {', '.join(v)}")
        return
    desconocidos = [n for n in a.nombres if n not in TE.TEMAS and n not in DECKS]
    if desconocidos:
        raise SystemExit(f"no reconozco: {desconocidos}. Temas: {', '.join(TE.TEMAS)}; presentaciones: {', '.join(DECKS)}")
    ids = [n for n in a.nombres if n in TE.TEMAS] or (list(TE.TEMAS) if a.todos else [TE.TEMA_OFICIAL])
    decks = [n for n in a.nombres if n in DECKS] or list(DECKS)
    decks = [k for k in decks if idioma in DECKS[k]]
    if not decks:
        raise SystemExit(f"ninguna de esas presentaciones tiene versión en {idioma!r}")
    mezcla = [m for m in (a.mezcla or "").split(",") if m]
    if mezcla:
        for m in mezcla:
            TE.obtener(m)
    for k in decks:
        cfg, ds = presentacion(k, idioma)
        faltan = [d.get("titulo") or d.get("frase") for d in ds if d["tipo"] in ("texto", "cita", "respaldo") and not d.get("sticker")]
        if faltan:
            print(f"[{k}] sin sticker:", faltan)
    preparar_videos([TE.obtener(i) for i in (mezcla or ids)], decks, idioma)
    if a.solo_videos:
        return
    copiar_fuentes()
    if mezcla:
        print(f"== Mezcla {' + '.join(mezcla)} (por {a.mezcla_por})", flush=True)
        for k in decks:
            construir(mezcla, k, idioma, a.mezcla_por)
        return
    for i in ids:
        print(f"== {TE.obtener(i).nombre}", flush=True)
        for k in decks:
            construir(i, k, idioma)


if __name__ == "__main__":
    main(sys.argv[1:])
