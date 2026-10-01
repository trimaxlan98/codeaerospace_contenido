"""Ponencia de divulgación «Redes Orbitales: el siguiente salto de la economía espacial».

~30 min, dos voces: Alan Rosas Palacios (red propia de satélites con IA 6G) y
Yuritzi Elena Ordaz Huerta (seguidor satelital inteligente que la red necesita).

NO renderiza nada: incrusta los mp4 ya existentes y les pone título y subtítulo en la diapositiva.
El guión completo va en las notas de cada diapositiva y en exports/GUION_REDES_ORBITALES.md
(con la hora estimada de cada diapositiva, calculada a RITMO palabras/min).
Salidas: exports/presentaciones/redes_orbitales_divulgacion_{oscuro,claro}.pptx
No borra ni pisa las otras presentaciones.
Uso: python3 ponencia_divulgacion.py
"""
import io
import sys
from pathlib import Path

from PIL import Image
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).parent))
from catalogo import BLOQUES  # noqa: E402
from empaquetar_ponencia import (ACENTO, EXP, H, TINTA, W, diapositiva, nombre_legible, nueva,  # noqa: E402
                                 poster, texto)
from tesis_decks import poster_de_video, video_final  # noqa: E402  (video sin el fundido final + póster del último cuadro)

RITMO = 140  # palabras por minuto al hablar en divulgación
CARPETA = {p[0]: c for _, c, ps in BLOQUES for p in ps}
MUESTRA = {p[0]: p[1] for _, _, ps in BLOQUES for p in ps}
TITULO = "Redes Orbitales"
LEMA = "El siguiente salto de la economía espacial"
A, Y, AY = "Alan", "Yuritzi", "Alan y Yuritzi"

# Tipos: portada · texto · seccion · video · cierre
# portada: (tipo, quien, guion)
# texto  : (tipo, quien, titulo, subtitulo, [viñetas], sticker|None, guion)
# seccion: (tipo, quien, titulo, subtitulo, guion)
# video  : (tipo, quien, pieza, titulo, subtitulo, guion)
DIAPOS = [
    ("portada", AY,
     "Buenas tardes. Somos Alan Rosas Palacios y Yuritzi Elena Ordaz Huerta, y venimos de Co.De Aerospace. "
     "Hoy queremos contarles algo que ya está pasando sobre nuestras cabezas: el espacio se está convirtiendo en una red. "
     "No en un lugar al que se va de visita, sino en infraestructura, como las carreteras o la fibra óptica. "
     "Y cuando una infraestructura nueva aparece, la economía suele dar un salto. Vamos a ver por qué, y qué hace falta para que ocurra. "
     "Antes de empezar, una pregunta para ustedes: ¿cuántas veces, hoy, usaron algo que depende de un satélite sin darse cuenta? "
     "El mapa del teléfono, la hora exacta, el pronóstico del clima, un pago con tarjeta. Casi todos, casi siempre. Esa es la infraestructura invisible de la que vamos a hablar."),

    ("texto", AY, "Dos tesis, una misma red", "Quiénes somos y qué construimos",
     ["Alan: una red propia de satélites gobernada con inteligencia artificial 6G",
      "Yuritzi: el seguidor satelital inteligente que esa red necesita para conectarse",
      "Dos tesis distintas que se necesitan la una a la otra"], "CubeSatDespiece",
     "Estamos trabajando en dos tesis distintas sobre tecnología satelital. "
     "Yo, Alan, trabajo en la red: cómo se organiza una red propia de satélites y cómo puede gobernarse con inteligencia artificial, "
     "en el marco de la sexta generación de redes móviles, el 6G. "
     "Y yo, Yuritzi, trabajo en el seguidor satelital inteligente: la estación en tierra que apunta la antena y sigue al satélite sin perderlo. "
     "Sin ese seguimiento, la mejor red del mundo no tiene con qué conectarse a tierra. "
     "Por eso decimos que son dos tesis, pero una sola red. La charla sigue ese camino: de la física básica, a la red, a la inteligencia, y de vuelta a la antena."),

    ("texto", AY, "Ruta de la charla", "Treinta minutos, de la física a la economía",
     ["1 · El espacio ya es infraestructura",
      "2 · De un satélite a una red",
      "3 · Tierra, aire y espacio: una sola red",
      "4 · Una red que se gobierna sola",
      "5 · Confianza: validar antes de creer",
      "6 · Apuntar bien: el seguidor satelital",
      "7 · Probar en tierra · 8 · Nuevas fronteras"], None,
     "Esta es la ruta. Empezamos con lo esencial: por qué un satélite se queda arriba. "
     "Después pasamos de un satélite a muchos, y de muchos a una red que integra tierra, aire y espacio. "
     "Luego viene la parte de inteligencia: cómo una red así puede gobernarse sola, y por qué no basta con decir que funciona, hay que demostrarlo. "
     "En la segunda mitad, Yuritzi nos lleva a la antena en tierra. Y cerramos con lo que sigue: probar antes de lanzar, nuevos servicios, y los límites que no podemos ignorar."),

    # ── 1
    ("seccion", Y, "El espacio ya es infraestructura", "Lo que la economía da por hecho hasta que falta",
     "Empecemos por lo básico. Cuando usamos un mapa en el teléfono o una videollamada en un barco, damos por hecho lo que hay arriba, "
     "igual que damos por hecho el agua del grifo. Veamos qué mantiene todo eso allá arriba."),
    ("video", Y, "CaidaLibreNewton", "Un satélite no flota: cae y sigue cayendo", "El cañón de Newton",
     "Newton lo imaginó con un cañón en una montaña muy alta. Si disparamos despacio, la bala cae cerca. Si disparamos más rápido, cae más lejos. "
     "Y si disparamos lo bastante rápido, la Tierra se curva por debajo de la bala tan deprisa como ella cae, y nunca llega al suelo. "
     "Eso es una órbita: caer alrededor de la Tierra. A esa altura hablamos de unos ocho kilómetros por segundo. "
     "Todo lo que sigue en esta charla descansa en esta idea sencilla."),
    ("video", Y, "OrbitasLEOMEOGEO", "Cada altura sirve para algo distinto", "Órbitas baja, media y geoestacionaria",
     "Según la altura, un satélite tiene un periodo distinto. Las órbitas bajas, LEO, dan una vuelta en aproximadamente hora y media. "
     "Las medias, MEO, tardan varias horas. Y a unos treinta y seis mil kilómetros, un satélite tarda exactamente lo que tarda la Tierra en girar, "
     "y parece quieto en el cielo: es la órbita geoestacionaria. "
     "Cada altura cambia tres cosas que le importan a un negocio: cuánto suelo cubre, cuánto tarda la señal, y cuánto cuesta llegar hasta ahí."),
    ("video", Y, "HuellaCobertura", "Más cerca, pero ve menos suelo", "El precio de la cercanía es la cantidad",
     "Desde una órbita alta, un solo satélite ve una gran parte del planeta. Desde una baja, ve un círculo mucho más pequeño. "
     "Pero estar cerca tiene ventajas enormes en velocidad de respuesta. "
     "El costo de esa cercanía es que se necesitan muchos satélites para cubrir todo. "
     "Ese detalle geométrico es el origen de las constelaciones, y de gran parte de la economía que viene."),
    ("video", Y, "TrazaTerrestre", "Así se ve un satélite desde el suelo", "Cada pasada es una oportunidad de conectar",
     "Desde el suelo, un satélite bajo no se queda: aparece, cruza el cielo en unos minutos y se va. "
     "Este dibujo es su rastro sobre el mapa, vuelta tras vuelta. "
     "Cada pasada es una ventana para conectar, y la antena tiene que estar lista para ella, en el lugar y el momento exactos. "
     "Guarden esta imagen: más adelante volveremos a ella. Ahora Alan nos cuenta qué pasa cuando no es un satélite, sino cientos."),

    # ── 2
    ("seccion", A, "De un satélite a una red", "El salto no es el satélite: es el servicio",
     "Aquí está el cambio de fondo. Un satélite es un aparato. Una red de satélites es un servicio. "
     "Y los servicios son los que mueven la economía."),
    ("video", A, "WalkerDelta3D", "Una constelación es un sistema, no una flota", "Geometría pensada para cubrir el planeta",
     "Una constelación no es una nube de satélites al azar. Tiene una geometría diseñada. "
     "Lo que ven es un patrón clásico, llamado Walker Delta: varios planos orbitales inclinados, con satélites repartidos con la misma separación. "
     "Es un ejemplo de diseño, no una constelación comercial en particular. "
     "La idea es que siempre haya algún satélite disponible sobre cada punto de la Tierra."),
    ("video", A, "PlanosOrbitales", "La cobertura se construye satélite a satélite", "Cuántos hacen falta y qué cuesta cubrir",
     "Con un plano orbital, cubrimos una franja. Con cada plano nuevo, la cobertura crece. "
     "Aquí está la primera pregunta de negocio: ¿cuántos satélites necesito para cubrir la zona que me interesa? "
     "Es una pregunta de ingeniería, pero su respuesta define la inversión inicial, y con ella, quién puede y quién no puede entrar a este mercado."),
    ("video", A, "ComparaLatencia", "La distancia se paga en tiempo", "Solo propagación: baja, media y alta",
     "La luz viaja rápido, pero el espacio es grande. Esta comparación muestra solo el tiempo que tarda la señal en ir y venir, sin contar equipos. "
     "En una órbita baja hablamos de unos pocos milisegundos. En la geoestacionaria, de casi un cuarto de segundo, solo de ida y vuelta. "
     "Para leer un correo, da igual. Para una videollamada, una operación financiera o un vehículo remoto, cambia todo. "
     "Por eso las órbitas bajas abren mercados que antes no estaban disponibles."),
    ("video", A, "HandoverSatelital", "El usuario no ve satélites, ve continuidad", "El relevo entre satélites",
     "Un satélite bajo se va en minutos. Para que el usuario no lo note, la conexión salta al siguiente satélite antes de perder el anterior. "
     "A eso se le llama relevo, o handover. "
     "El valor comercial no está en el hardware que pasa por arriba, sino en que la conexión no se corta. "
     "Nadie paga por tener un satélite: se paga por un servicio que sigue ahí."),
    ("video", A, "MallaISL", "Los satélites también se conectan entre sí", "Enlaces intersatélite: una red en el vacío",
     "Además de hablar con la tierra, los satélites pueden hablar entre ellos, por enlaces intersatélite. "
     "Así se forma una malla que lleva la información de un continente a otro sin bajar a tierra en cada salto. "
     "Es una topología nueva, y con ella aparecen modelos de negocio nuevos: el espacio no solo como antena, sino como red de transporte."),

    # ── 3
    ("seccion", A, "Tierra, aire y espacio: una sola red", "La visión del 6G",
     "La sexta generación de redes móviles plantea algo ambicioso: que el usuario no sepa, ni tenga que saber, por dónde le llega la señal."),
    ("video", A, "ZoologicoOrbital", "Cada plataforma cubre un hueco distinto", "Plataformas de gran altitud, LEO, MEO y GEO",
     "Este es el zoológico de plataformas. Globos y aeronaves a gran altura, satélites bajos, medios y geoestacionarios. "
     "Cada uno tiene su retardo, su cobertura y su costo. "
     "No compiten entre sí: se complementan. Las redes del futuro combinan las capas según lo que cada aplicación necesita."),
    ("video", A, "EspacioAireTierra", "Tres niveles, un solo servicio", "Una red integrada",
     "Aquí lo vemos integrado. Abajo, las antenas de siempre. En medio, plataformas a gran altura. Arriba, los satélites. "
     "El usuario tiene un solo teléfono y un solo contrato, y la red decide por dónde llevar cada comunicación. "
     "Esa decisión es lo que vuelve la red difícil de operar, y es justo donde entra la inteligencia artificial."),
    ("video", A, "ArquitecturaNTN", "El satélite entra al estándar móvil", "Redes no terrestres en el estándar 3GPP",
     "Esto no es ciencia ficción: el satélite ya entró al estándar de telefonía móvil, bajo el nombre de redes no terrestres, o NTN. "
     "Un teléfono, un satélite, una estación de enlace, el núcleo de la red, e Internet. "
     "Los estándares importan mucho en economía: cuando todos hablan el mismo idioma, los mercados se abren y los equipos se abaratan."),
    ("video", A, "Interferencia", "Compartir el cielo exige coordinación", "Interferencia entre haces y constelaciones",
     "Hay un detalle que rara vez sale en las noticias: el cielo se comparte. "
     "Cuando dos haces se traslapan sobre la misma zona y la misma frecuencia, se estorban. Eso es interferencia. "
     "Con pocas constelaciones se resuelve con acuerdos y buena ingeniería, pero con muchas, la coordinación se vuelve un tema económico y político. "
     "El espectro es un recurso limitado y valioso, igual que un terreno."),
    ("texto", A, "¿De dónde sale el valor?", "La economía detrás de la red",
     ["Conectividad donde la fibra no llega: campo, minería, mar, aire",
      "Servicios sensibles al tiempo: logística, sincronía, emergencias",
      "Datos desde el espacio, procesados en el borde, entregados como decisión",
      "Operación automática: la red que se valida y se gobierna sola cuesta menos"], "Tierra3DSatelites",
     "Entonces, ¿de dónde sale el valor? Primero, de llevar conectividad donde la fibra nunca va a llegar: el campo, la minería, el mar, la aviación. "
     "Segundo, de los servicios que dependen del tiempo, como la logística o la respuesta a emergencias. "
     "Tercero, de los datos que se generan en el espacio y se convierten en decisiones. "
     "Y cuarto, del ahorro de operar la red de forma automática. Ese último punto es el que sigue."),

    ("texto", A, "Tres casos, tres economías", "Lo que cambia cuando la conectividad llega a todas partes",
     ["Campo y minería: monitoreo y maquinaria conectados donde no hay antenas",
      "Mar y aire: rutas, seguridad y servicios a bordo sin zonas muertas",
      "Emergencias: comunicación cuando la red terrestre cae"], "HuellaCobertura",
     "Pongamos tres casos concretos. Primero, el campo y la minería: sensores, maquinaria y monitoreo en lugares donde nunca habrá una antena celular. "
     "Segundo, el mar y el aire: barcos y aviones que hoy pasan horas sin buena conexión y podrían tener servicio continuo. "
     "Y tercero, las emergencias: cuando un huracán o un sismo derrumban la red terrestre, la red del cielo sigue en pie. "
     "Ninguno de estos casos es hipotético; lo que cambia es que la conectividad deja de depender de dónde estás."),

    # ── 4
    ("seccion", A, "Una red que se gobierna sola", "Con miles de nodos, operar a mano no escala",
     "Una constelación grande es una red que nunca se queda quieta. Y una red que se mueve tanto necesita un cerebro."),
    ("video", A, "TopologiaRespira", "La red cambia todo el tiempo", "Enlaces que aparecen y desaparecen",
     "Miren cómo respira esta red. Los enlaces aparecen cuando dos satélites se ven, y desaparecen cuando se pierden de vista. "
     "En una red terrestre, los cables se quedan donde están. Aquí, la forma de la red cambia cada minuto. "
     "Un operador humano no puede reconfigurarla a la velocidad a la que se mueve. Hace falta que la propia red aprenda a decidir."),
    ("video", A, "CicloPADA", "Percibir, analizar, decidir, actuar", "La arquitectura PADA",
     "En mi tesis, esa idea se organiza en cuatro pasos: Percepción, Análisis, Decisión y Acción. Se llama PADA. "
     "La red mide lo que pasa, entiende qué significa, decide qué hacer y lo hace, y vuelve a medir. "
     "Es un ciclo continuo, parecido a como nosotros manejamos un coche."),
    ("video", A, "MuchosAgentesCTDE", "Se entrena en conjunto, se decide en cada nodo", "Aprendizaje por refuerzo multiagente",
     "¿Cómo aprende una red así? Con aprendizaje por refuerzo multiagente. "
     "Durante el entrenamiento, un mezclador ve el panorama completo y enseña a los agentes a colaborar. "
     "Pero cuando la red ya opera, cada nodo decide por su cuenta, con lo que ve. "
     "Es entrenamiento centralizado, ejecución descentralizada. En nuestras pruebas usamos tres agentes: dos satélites y un gateway."),
    ("video", A, "PoliticaEnrutaTrafico", "Una política que se adapta evita el cuello de botella", "Ilustrativo: estática frente a adaptativa",
     "Un ejemplo ilustrativo, sin cifras. Con una política fija, cuando una ruta se congestiona, el tráfico sigue llegando y el retardo sube. "
     "Con una política adaptativa, la red cambia de ruta a tiempo. "
     "Este dibujo es una idea, no un resultado medido. Los resultados medidos los veremos en un momento, y con mucho cuidado."),
    ("video", A, "GrafoGNN", "La red como grafo: los nodos se cuentan lo que ven", "Redes neuronales de grafos",
     "Una red de satélites se puede dibujar como un grafo: los satélites son puntos y los enlaces son líneas. "
     "Existe una familia de modelos, las redes neuronales de grafos, hechas justo para esto. "
     "Cada nodo le pregunta a sus vecinos qué ven, combina las respuestas, y así, poco a poco, la información se esparce por toda la red. "
     "Lo interesante es que funcionan aunque el grafo cambie de forma, y la nuestra cambia todo el tiempo. Es una de las herramientas que exploramos."),
    ("video", A, "IAaBordo", "Procesar en órbita es enviar menos a tierra", "Inteligencia artificial en el borde",
     "La inteligencia también puede viajar en el satélite. Si el satélite detecta a bordo lo que importa, no hace falta bajar todo a tierra. "
     "Es enviar menos datos, y por eso liberar ancho de banda, que en el espacio es un recurso caro. "
     "Es la misma lógica que un buen editor: no mandas toda la grabación, mandas lo que sirve."),

    # ── 5
    ("seccion", A, "Confianza: validar antes de creer", "Una economía autónoma necesita evidencia",
     "Aquí quiero ser honesto. Es fácil decir que una inteligencia artificial mejora una red. Lo difícil es demostrarlo sin engañarse."),
    ("video", A, "SateliteMiente", "¿Y si un nodo miente?", "Consenso tolerante a fallos bizantinos",
     "Una red autónoma tiene otro riesgo: que uno de los nodos falle, o peor, que mienta. "
     "Existen métodos de consenso, como PBFT, que toleran nodos deshonestos si hay suficientes honestos: n mayor o igual a tres f más uno. "
     "Aquí el ejemplo usa siete nodos y dos defectuosos, y la red los aísla. Es una ilustración del principio, no un resultado de nuestras pruebas."),
    ("video", A, "MargenAdaptativo", "Primero: ¿el banco de pruebas puede distinguir a la IA?", "Margen Adaptativo: una propiedad del banco",
     "Esta es la idea más importante de mi trabajo. Antes de evaluar una inteligencia artificial, hay que preguntar si el banco de pruebas es capaz de distinguirla. "
     "Si el entorno es tan sencillo que una política fija ya es casi óptima, no hay nada que aprender, y cualquier mejora aparente es ruido. "
     "Lo medimos con el Margen Adaptativo: la diferencia entre un oráculo que lo sabe todo y la mejor política estática. "
     "En nuestra primera versión del entorno, era del uno punto siete por ciento: un banco que no discriminaba. En la versión corregida, treinta y uno punto ocho. "
     "Y un detalle que nos gusta contar: una heurística clásica quedó por debajo de la mejor política fija. La intuición no siempre gana."),
    ("video", A, "CompuertasValidacion", "Cada compuerta se gana", "De G0 a G3: una falla real y su corrección",
     "Organizamos la validación en compuertas. G0 es un control negativo: comprobar que el banco no ve mejoras donde no las hay. "
     "G1 verifica que el margen sea suficiente. G2 compara el aprendizaje contra políticas fijas, y en él tuvimos una falla real en el primer intento, que corregimos. "
     "En simulación, con tres agentes, el aprendizaje mejoró entre nueve y dieciséis por ciento. "
     "Y la compuerta tres todavía no la hemos lanzado: está pendiente. Decirlo también es parte de la honestidad científica. Una compuerta pendiente no es un fracaso; es una afirmación que aún no nos hemos ganado el derecho de hacer."),

    ("texto", A, "Lo que sabemos y lo que falta", "Decir el alcance de cada resultado",
     ["Validado en simulación: compuertas G0, G1, G2a y G2b",
      "Pendiente: la compuerta G3, todavía sin lanzar",
      "Los nodos maliciosos son un ejemplo ilustrativo, no un resultado",
      "Lo ilustrativo se marca como ilustrativo"], "MargenAdaptativo",
     "Quiero cerrar esta parte con un balance. Lo que ya sabemos: en simulación, y con tres agentes, pasamos las compuertas G0, G1, G2a y G2b. "
     "Lo que falta: la compuerta tres. Los nodos que mienten son un ejemplo para explicar una idea, no un resultado nuestro. "
     "Y cuando una figura es ilustrativa, la marcamos como ilustrativa. "
     "Creemos que, si vamos a proponer que las redes se gobiernen solas, lo mínimo es ser claros sobre qué sabemos y qué no."),

    # ── 6
    ("seccion", Y, "Apuntar bien: el seguidor satelital inteligente", "Sin antena que lo siga, la red no llega a tierra",
     "Gracias, Alan. Volvamos a tierra. Todo lo que acabamos de ver depende de algo que parece simple: que una antena apunte al satélite correcto, en el momento correcto, y siga apuntando mientras pasa."),
    ("video", Y, "VistaPolarPase", "Un pase visto desde la antena", "Sale, culmina y se oculta",
     "Este es el mismo pase, pero visto desde la antena. El centro es el cielo justo arriba, y el borde es el horizonte. "
     "El satélite sale por un lado, sube hasta su punto más alto, y se oculta por el otro. "
     "Tres momentos importan: cuando aparece, cuando culmina, y cuando desaparece. "
     "La antena debe estar apuntando antes de que aparezca, y acompañar todo el recorrido."),
    ("video", Y, "CadenaTLEaAntena", "Del dato orbital al movimiento de la antena", "TLE, SGP4, coordenadas y motores",
     "¿Cómo sabe la antena a dónde mirar? Empezamos con un dato público llamado TLE, que describe la órbita del satélite. "
     "Un modelo llamado SGP4 lo convierte en una posición en el espacio. "
     "Luego esa posición se traduce a la vista de quien está en tierra: acimut y elevación, es decir, hacia dónde girar y cuánto levantar. "
     "El controlador transforma esos ángulos en órdenes para los motores. Es una cadena larga, y cada eslabón tiene su margen de error."),
    ("video", Y, "MontajeAzEl", "Dos ejes para cubrir todo el cielo", "La montura de acimut y elevación",
     "Para apuntar a cualquier punto del cielo basta con dos giros. Uno horizontal, el acimut, que barre alrededor. Y uno vertical, la elevación, que levanta el plato. "
     "Es la misma idea de un telescopio o de una cámara de vigilancia. "
     "Es una solución sencilla y barata, y por eso muy usada. Pero tiene una limitación que veremos en unos minutos."),
    ("video", Y, "AntenaSiguiendo", "La antena sigue al satélite", "Seguimiento sobre la máscara de elevación",
     "Aquí se ve el resultado: el plato girando para seguir al satélite. "
     "La zona más cercana al horizonte se descarta con una máscara de elevación, porque allí hay edificios, árboles y mucha atmósfera. "
     "Solo se sigue al satélite cuando vale la pena. Es una decisión pequeña que ahorra mucha energía y evita malas conexiones."),
    ("video", Y, "LazoPID", "Apuntar es un problema de control", "La antena persigue la referencia",
     "Seguir a un satélite es un problema de control. La referencia se mueve, y la antena tiene inercia, fricción y retardo. "
     "Un lazo de control corrige continuamente la diferencia entre dónde debería estar y dónde está. "
     "Ajustarlo bien es la diferencia entre una antena que sigue con suavidad, y una que oscila, se pasa o llega tarde."),
    ("video", Y, "Keyhole", "El pase que las monturas temen", "Cuando el satélite pasa por el cenit",
     "Hay un caso especial que da dolores de cabeza. Cuando el satélite pasa casi justo encima, el acimut tiene que girar muy rápido, más rápido de lo que el motor puede. "
     "Se le llama efecto ojo de cerradura, o keyhole. "
     "Es una limitación conocida de las monturas de acimut y elevación, y es uno de los puntos donde un seguidor inteligente marca la diferencia frente a uno ingenuo. "
     "[YURITZI: aquí explicar la estrategia propia para el keyhole y mencionar el resultado medido, si lo hay.]"),
    ("video", Y, "DopplerEnS", "La frecuencia también se mueve", "Compensación Doppler durante el pase",
     "Además de moverse en el cielo, el satélite mueve su frecuencia. Al acercarse, la señal se percibe más aguda. Al alejarse, más grave. Es el efecto Doppler. "
     "Si el receptor no lo compensa, deja de escuchar al satélite aunque la antena apunte perfecto. "
     "Por eso el seguidor no solo mueve motores: también sintoniza el receptor en cada instante del pase."),
    ("video", Y, "NochePases", "Una noche, varios pases", "Planear las ventanas útiles",
     "Un satélite bajo pasa varias veces por noche, y cada pase dura pocos minutos. Algunos son muy bajos, otros pasan casi por encima. "
     "Entonces la estación no solo apunta: también decide a cuáles pases vale la pena conectarse. "
     "Planear la noche es parte del trabajo. Una buena planeación multiplica la cantidad de datos que se pueden bajar con el mismo equipo."),
    ("video", Y, "GemeloDigitalATP", "Antena real y gemelo digital, sincronizados", "Del seguimiento al banco de pruebas",
     "Por último, un gemelo digital: una copia virtual de la antena que se mueve igual que la real. "
     "Con él podemos probar cambios, ensayar pases y detectar diferencias sin arriesgar el equipo. "
     "Y es el puente hacia lo que sigue: el banco de pruebas donde la red de Alan y el seguidor se encuentran."),

    ("texto", Y, "Qué aporta el seguidor a la red", "La otra mitad del servicio",
     ["Un enlace fiable en cada pase",
      "Menos intervención humana en la operación",
      "Un gemelo digital para ensayar sin riesgo",
      "Una base para operar una red propia"], "AntenaSiguiendo",
     "Resumiendo mi parte: el seguidor le da a la red un enlace confiable en cada pase, reduce la intervención humana, y con el gemelo digital permite ensayar sin riesgo. "
     "Es, en pocas palabras, lo que hace posible operar una red propia desde tierra. "
     "[YURITZI: aquí añadir el resultado propio de la tesis, con su alcance.] "
     "Y ahora, juntos, veamos cómo se prueba todo esto antes de lanzar."),

    # ── 7
    ("seccion", AY, "Probar en tierra antes de lanzar", "Cada hora de banco es un riesgo menos en órbita",
     "Lanzar cuesta caro, y arreglar algo en órbita es casi imposible. Por eso, en la industria espacial, probar en tierra no es un lujo: es la manera de ahorrar."),
    ("video", A, "CubeSatDespiece", "El satélite bajo prueba", "Un CubeSat 3U en vista explotada",
     "Este es un CubeSat de tres unidades: más o menos el tamaño de una caja de zapatos. "
     "Estructura, baterías, computadora, radio, paneles. "
     "Estas plataformas hicieron accesible el espacio a universidades y empresas pequeñas, y son un buen ejemplo de qué se puede probar antes de ir arriba."),
    ("video", A, "BancoHardwareEnLaLazo", "El hardware real, en un cielo simulado", "Banco de pruebas con hardware en el lazo",
     "En el banco de pruebas conectamos el hardware real a un simulador del cielo y a un emulador del canal. "
     "El equipo cree que está en órbita, pero sigue en el laboratorio. "
     "Es aquí donde se encuentran nuestras dos tesis: la red con su inteligencia, y el seguidor con su antena."),
    ("video", Y, "EmuladorCanal", "El canal espacial, en la mesa del laboratorio", "Retardo, Doppler, atenuación y ruido",
     "El emulador de canal recrea lo que le pasa a la señal en el viaje: llega tarde, con la frecuencia movida, más débil y con ruido. "
     "Los valores del ejemplo son estándar, no mediciones de nuestro banco. "
     "Con esto, probamos en un escritorio lo que en el espacio sería muy difícil de repetir."),

    # ── 8
    ("seccion", AY, "Nuevas fronteras y sus límites", "Nuevos mercados, y el costo de no cuidar el recurso",
     "Cerramos con dos miradas: lo que se abre, y lo que nos toca cuidar."),
    ("video", A, "ServicioEnOrbita", "El satélite deja de ser desechable", "Reparar, reabastecer y extender la vida útil",
     "Hoy, cuando un satélite se queda sin combustible o falla, se pierde. Un satélite servidor puede acercarse, acoplarse y darle una segunda vida. "
     "Eso abre una línea de servicios completa: mantenimiento, reabastecimiento, remolque al final de su vida. "
     "Es como pasar de vehículos desechables a vehículos con taller."),
    ("video", Y, "BasuraEspacial", "Crecer sin ensuciar el camino", "El efecto Kessler",
     "Pero hay un límite físico. Cada satélite y cada fragmento en órbita es un posible choque. "
     "Cuando hay demasiados, un choque genera fragmentos que provocan más choques: es el efecto Kessler. "
     "Una economía orbital que no cuida sus órbitas se acaba a sí misma. Por eso la sostenibilidad no es un adorno: es parte del modelo."),

    # ── cierre
    ("texto", Y, "Qué necesita para ocurrir", "Cuatro condiciones para el siguiente salto",
     ["Estándares abiertos que integren tierra y espacio",
      "Espectro coordinado y órbitas sostenibles",
      "Autonomía verificable: evidencia antes de desplegar",
      "Talento y bancos de prueba propios"], "PlanosOrbitales",
     "Para que este salto ocurra hacen falta cuatro cosas. Estándares abiertos, para que la tierra y el espacio se integren. "
     "Un uso ordenado del espectro y de las órbitas. Autonomía verificable: si una red se gobierna sola, tiene que demostrar por qué confiamos en ella. "
     "Y talento con lugares donde probar. Esto último es lo que más nos importa a nosotros."),
    ("texto", AY, "Dónde entramos nosotros", "Co.De Aerospace",
     ["Alan: una red satelital gobernada con IA, validada con evidencia",
      "Yuritzi: un seguidor satelital inteligente para conectarse a ella",
      "Un banco de pruebas donde ambas se encuentran",
      "El camino: del modelo al hardware"], "CubeSatDespiece",
     "Nosotros estamos en ese camino. Alan trabaja en la red y en cómo demostrar que una inteligencia artificial la mejora de verdad. "
     "Yuritzi, en el seguidor que permite que esa red se conecte con la tierra. "
     "Y compartimos un banco de pruebas donde ambos trabajos se encuentran, del modelo al hardware. "
     "Lo hacemos en Co.De Aerospace porque creemos que este conocimiento tiene que crecer aquí, con gente joven y con lugares para probar."),
    ("texto", AY, "Tres ideas para llevarse", "Si solo recuerdan tres cosas",
     ["El valor está en el servicio, no en el satélite",
      "Una red que cambia sola necesita autonomía",
      "La autonomía necesita evidencia"], "WalkerDelta3D",
     "Si solo se llevan tres cosas de esta charla, que sean estas. Una: el valor está en el servicio, no en el satélite. Nadie compra un aparato en órbita; compra una conexión que no se corta. "
     "Dos: una red que cambia todo el tiempo necesita autonomía, porque a mano no se puede. "
     "Y tres: la autonomía necesita evidencia. Sin pruebas, una red que se gobierna sola es una apuesta; con pruebas, es infraestructura."),
    ("cierre", AY,
     "El espacio ya no es solo un destino: es una red. Y las redes, cuando aparecen, cambian la economía. "
     "Muchas gracias por su atención. Con gusto respondemos sus preguntas."),
]


# Ampliaciones de guion por pieza (se añaden al final del texto de la diapositiva de video).
AMPLIA = {
    "CaidaLibreNewton": "Un ejemplo cotidiano: los astronautas de la estación espacial no flotan porque no haya gravedad. Flotan porque están en caída libre, junto con la estación, todo el tiempo. La gravedad allí es casi la misma que aquí.",
    "OrbitasLEOMEOGEO": "Una analogía: es como elegir el piso de un edificio. En el piso bajo, ves poco pero estás cerca de la calle. En la azotea, ves toda la ciudad pero estás lejos del ruido y de la gente.",
    "TrazaTerrestre": "Esa ondulación tiene explicación: mientras el satélite da su vuelta, la Tierra gira por debajo, y por eso cada órbita cae en un lugar distinto del mapa.",
    "ComparaLatencia": "Y fíjense en que aquí solo cuenta la distancia. En la vida real se suman el procesamiento y las esperas, pero la distancia es el piso que ninguna tecnología puede bajar.",
    "MallaISL": "Una comparación útil: es como la diferencia entre una carretera con un solo camino y una ciudad con muchas calles. Si una se bloquea, el tráfico toma otra. La malla en el cielo le da a la red esa capacidad de rodear un problema.",
    "ArquitecturaNTN": "Para quien usa el teléfono, nada cambia. Para quien construye la red, cambia todo: el satélite deja de ser un mundo aparte y pasa a ser una parte más de la misma red.",
    "CicloPADA": "Lo importante de PADA no es el nombre, sino la separación en pasos. Cuando algo falla, sabemos en cuál: si la red no vio el problema, si lo entendió mal, si decidió mal, o si no pudo ejecutar.",
    "MuchosAgentesCTDE": "Es como entrenar a un equipo de fútbol: en los entrenamientos el técnico ve toda la cancha y da instrucciones; en el partido, cada jugador decide con lo que tiene enfrente.",
    "IAaBordo": "Y esto es aún más valioso en el espacio: cada bit que baja a tierra usa un enlace escaso y caro, así que decidir a bordo qué vale la pena enviar es una ventaja económica directa.",
    "MargenAdaptativo": "Piénsenlo así: si todos los caminos llevan al mismo lugar, no tiene sentido evaluar quién es mejor guía. Primero hay que construir un terreno donde elegir bien importe. Esa fue la corrección más importante que hicimos.",
    "CompuertasValidacion": "Este esquema es incómodo a propósito: nos obliga a detenernos cada vez que algo no está demostrado. Preferimos tardar más y llegar con resultados que podamos defender.",
    "VistaPolarPase": "Si alguna vez miraron pasar la Estación Espacial Internacional, vieron exactamente esto: aparece por un lado, sube, y desaparece por el otro en pocos minutos.",
    "CadenaTLEaAntena": "Un detalle importante: el TLE envejece. Cuanto más antiguo es, menos precisa es la posición calculada, así que mantenerlo actualizado es parte de la operación diaria.",
    "AntenaSiguiendo": "Y en esta demostración se ve algo que a veces se olvida: la antena no es un robot que acierta a la primera; es un sistema que corrige continuamente su apuntamiento.",
    "LazoPID": "Un ejemplo doméstico: es como conducir un coche por una autopista recta con el volante. Nadie lo mantiene perfecto; se corrigen miles de pequeños desvíos, y se hace tan suave que no lo notamos.",
    "Keyhole": "Una analogía: es como intentar seguir con la vista un avión que pasa justo sobre tu cabeza. Al principio apenas mueves la cabeza, pero de pronto tienes que girar todo el cuerpo muy rápido.",
    "DopplerEnS": "Es el mismo efecto que percibimos con la sirena de una ambulancia: al acercarse suena más aguda, al alejarse más grave. En un satélite ocurre con las ondas de radio, y a velocidades mucho mayores.",
    "GemeloDigitalATP": "La lógica es la misma que se usa para probar aviones en simuladores antes de volar: equivocarse en la copia virtual sale gratis, en el equipo real sale caro.",
    "EmuladorCanal": "La ventaja para una universidad o una empresa pequeña es enorme: en vez de contratar un lanzamiento para saber si algo funciona, se ensaya en tierra, tantas veces como haga falta.",
    "BasuraEspacial": "La buena noticia es que esto se puede prevenir: diseñar satélites que salgan de órbita al final de su vida, y coordinar la operación entre todos los que compartimos el mismo espacio.",
    "ServicioEnOrbita": "Es un cambio parecido al de pasar de la botella desechable a la recargable: el mismo objeto, pero con una vida útil y un modelo de negocio completamente distintos.",
}
for _i, _d in enumerate(DIAPOS):
    if _d[0] == "video" and _d[2] in AMPLIA:
        DIAPOS[_i] = _d[:-1] + (_d[-1] + " " + AMPLIA[_d[2]],)


# ─────────────────────────── utilidades ───────────────────────────
def palabras(t):
    return len(t.split())


def minutos(d):
    return palabras(d[-1]) / RITMO


def sticker(tema, clase):
    src = EXP / "png" / "stickers" / tema / CARPETA[clase] / f"{clase}.png"
    im = Image.open(src).convert("RGBA")
    im.thumbnail((1800, 1800), Image.LANCZOS)
    b = io.BytesIO()
    im.save(b, "PNG", compress_level=6)
    b.seek(0)
    return b, im.size


def subtexto(s, tema, txt, x, y, w, tam=20):
    # subtítulo en el color de acento suave
    return texto(s, txt, x, y, w, Inches(0.6), tam, ACENTO[tema])


def etiqueta_voz(s, tema, quien):
    texto(s, quien, Inches(0.6), Inches(7.05), Inches(5), Inches(0.35), 12, TINTA[tema])


def notas(s, d, extra=""):
    s.notes_slide.notes_text_frame.text = f"[{d[1]}] {d[-1]}{extra}"


def hacer(prs, tema, d, n, num_seccion):
    tipo, quien = d[0], d[1]
    s = diapositiva(prs, tema)
    if tipo == "portada":
        texto(s, TITULO, Inches(0.9), Inches(1.9), Inches(11.5), Inches(1.6), 72, TINTA[tema], True)
        texto(s, LEMA, Inches(0.9), Inches(3.6), Inches(11.5), Inches(1.4), 34, ACENTO[tema])
        texto(s, "Alan Rosas Palacios  ·  Yuritzi Elena Ordaz Huerta", Inches(0.9), Inches(5.6), Inches(11.5),
              Inches(0.6), 24, TINTA[tema])
        texto(s, "Co.De Aerospace", Inches(0.9), Inches(6.2), Inches(6), Inches(0.6), 20, TINTA[tema])
        notas(s, d)
    elif tipo == "seccion":
        _, _, titulo, sub, _g = d
        texto(s, f"{num_seccion:02d}", Inches(0.9), Inches(2.0), Inches(4), Inches(1.2), 54, ACENTO[tema], True)
        texto(s, titulo, Inches(0.9), Inches(3.1), Inches(11.5), Inches(1.4), 42, TINTA[tema], True)
        subtexto(s, tema, sub, Inches(0.9), Inches(4.6), Inches(11.5), 26)
        etiqueta_voz(s, tema, quien)
        notas(s, d)
    elif tipo == "video":
        _, _, clase, titulo, sub, _g = d
        assert (EXP / CARPETA[clase] / tema / f"{clase}.mp4").exists(), clase
        mp4 = video_final(tema, clase, CARPETA[clase])
        texto(s, titulo, Inches(0.6), Inches(0.22), Inches(12.1), Inches(0.75), 30, TINTA[tema], True)
        subtexto(s, tema, sub, Inches(0.6), Inches(0.9), Inches(12.1), 20)
        vh = Inches(5.5)
        vw = int(vh * 16 / 9)
        s.shapes.add_movie(str(mp4), int((W - vw) / 2), Inches(1.5), vw, vh,
                           poster_frame_image=poster_de_video(tema, clase, CARPETA[clase]), mime_type="video/mp4")
        etiqueta_voz(s, tema, quien)
        notas(s, d, f"\n\n(Video: {nombre_legible(clase)} — {MUESTRA[clase]}. Clic para reproducir; al terminar queda el último cuadro.)")
    elif tipo == "texto":
        _, _, titulo, sub, puntos, clase, _g = d
        texto(s, titulo, Inches(0.7), Inches(0.4), Inches(12), Inches(0.9), 38, TINTA[tema], True)
        subtexto(s, tema, sub, Inches(0.7), Inches(1.25), Inches(12), 22)
        ancho = Inches(6.7) if clase else Inches(11.8)
        tb = s.shapes.add_textbox(Inches(0.7), Inches(2.1), ancho, Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, t in enumerate(puntos):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(16)
            r = p.add_run()
            r.text = "▪ " + t
            r.font.size, r.font.name = Pt(24), "Carlito"
            r.font.color.rgb = TINTA[tema]
        if clase:
            buf, (iw, ih) = sticker(tema, clase)
            esc = min(Inches(5.2) / iw, Inches(4.8) / ih)
            w, h = int(iw * esc), int(ih * esc)
            pic = s.shapes.add_picture(buf, Inches(7.7) + int((Inches(5.2) - w) / 2), Inches(2.0) + int((Inches(4.8) - h) / 2), w, h)
            pic._element.nvPicPr.cNvPr.set("descr", f"{nombre_legible(clase)}: {MUESTRA[clase]}")
        etiqueta_voz(s, tema, quien)
        notas(s, d)
    elif tipo == "cierre":
        texto(s, "El siguiente salto empieza en órbita", Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.4), 48, TINTA[tema], True)
        subtexto(s, tema, "Gracias · Preguntas", Inches(0.9), Inches(4.0), Inches(11.5), 30)
        texto(s, "Alan Rosas Palacios  ·  Yuritzi Elena Ordaz Huerta  ·  Co.De Aerospace", Inches(0.9), Inches(6.3),
              Inches(11.5), Inches(0.6), 18, TINTA[tema])
        notas(s, d)
    return s


def guion_md():
    tot = sum(minutos(d) for d in DIAPOS)
    L = [f"# Guion — {TITULO}: {LEMA}", "",
         f"Voces: **Alan Rosas Palacios** (red propia de satélites con IA 6G) y **Yuritzi Elena Ordaz Huerta** "
         f"(seguidor satelital inteligente). Ritmo: {RITMO} palabras/min → **{tot:.1f} min** de guion "
         f"({sum(palabras(d[-1]) for d in DIAPOS)} palabras), más el tiempo de las preguntas.", "",
         "Los videos ya existen y no se vuelven a renderizar; en PowerPoint arrancan con clic. "
         "Cada video dura 13–27 s: se puede seguir hablando sobre el último cuadro.", "",
         "> Las líneas `[YURITZI: …]` son huecos para que complete ella misma con su resultado medido.", "",
         "| # | Hora | Voz | Diapositiva |", "|---|---|---|---|"]
    t = 0.0
    filas, cuerpo = [], []
    for n, d in enumerate(DIAPOS, 1):
        titulo = d[2] if d[0] in ("seccion", "texto") else d[3] if d[0] == "video" else "Portada" if d[0] == "portada" else "Cierre"
        sub = d[3] if d[0] in ("seccion", "texto") else d[4] if d[0] == "video" else ""
        h = f"{int(t)}:{int((t % 1) * 60):02d}"
        filas.append(f"| {n} | {h} | {d[1]} | {titulo} |")
        cuerpo += [f"## {n}. {titulo}", f"*{sub}*  ·  **{d[1]}**  ·  {h}" + (f"  ·  video `{d[2]}`" if d[0] == "video" else ""),
                   "", d[-1], ""]
        t += minutos(d)
    return "\n".join(L + filas + [""] + cuerpo)


def construir(tema):
    prs = nueva(tema)
    sec = 0
    for n, d in enumerate(DIAPOS, 1):
        if d[0] == "seccion":
            sec += 1
        hacer(prs, tema, d, n, sec)
    destino = EXP / "presentaciones" / f"redes_orbitales_divulgacion_{tema}.pptx"
    prs.save(destino)
    print(destino.name, len(prs.slides), "diapositivas", f"{destino.stat().st_size / 1e6:.0f} MB")


if __name__ == "__main__":
    tot = sum(minutos(d) for d in DIAPOS)
    print(f"{len(DIAPOS)} diapositivas · {sum(palabras(d[-1]) for d in DIAPOS)} palabras · {tot:.1f} min a {RITMO} ppm")
    for q in (A, Y, AY):
        print(f"  {q}: {sum(minutos(d) for d in DIAPOS if d[1] == q):.1f} min")
    (EXP / "GUION_REDES_ORBITALES.md").write_text(guion_md(), encoding="utf-8")
    if "--solo-guion" not in sys.argv:
        for t in ("oscuro", "claro"):
            construir(t)
